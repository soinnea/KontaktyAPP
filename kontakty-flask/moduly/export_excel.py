from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
import os

BARVA_HLAVICKY = "1976D2"

def _hlavicka(ws, sloupce):
    ws.append(sloupce)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BARVA_HLAVICKY)
        cell.alignment = Alignment(horizontal="center")

def export_kontakty_excel(kontakty):
    wb = Workbook()
    ws = wb.active
    ws.title = "Kontakty"
    _hlavicka(ws, ['Jméno', 'Příjmení', 'Email', 'Telefon', 'Město', 'Pozice', 'Firma'])
    for k in kontakty:
        ws.append([
            k.get('Jmeno', ''), k.get('Prijmeni', ''), k.get('Email', ''),
            k.get('Telefon', ''), k.get('Mesto', ''), k.get('Pozice', ''),
            k.get('FirmaNazev', 'Solo')
        ])
    nazev = f"kontakty_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    cesta = os.path.join("vystupy", nazev)
    wb.save(cesta)
    return cesta

def export_firmy_excel(firmy):
    wb = Workbook()
    ws = wb.active
    ws.title = "Firmy"
    _hlavicka(ws, ['Název', 'Email', 'Telefon', 'Město', 'IČO', 'Segment'])
    for f in firmy:
        ws.append([
            f.get('Nazev', ''), f.get('Email', ''), f.get('Telefon', ''),
            f.get('Mesto', ''), f.get('Ico', ''), f.get('Segment', '')
        ])
    nazev = f"firmy_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    cesta = os.path.join("vystupy", nazev)
    wb.save(cesta)
    return cesta
